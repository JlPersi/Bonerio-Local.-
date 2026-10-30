"""Arma data/pvs.json, data/pvs_emisores.json y (opcional) data/pvs_cierre.json para la sección de Bonos provinciales.

Entrada: data/pvs_flujos.xlsx (flujos por 100 VN de cada bono subsoberano, hoja «Cashflow de Bonos»).
Opcional: SERIES=series.xlsx (hoja «Serie  de Bond»: cierres y monto por día) y YIELDS=yield-bonds.csv (TIR del cierre)
para la foto de cierres. Salida por bono (el ticker en pesos): tipo, emisión, vencimiento, tasa, residual, flujos
que faltan cobrar y las especies en dólares (d = MEP, c = cable), que en los provinciales no siguen la regla ticker + D/C.
"""
import json, os, sys, csv, datetime, statistics
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(HERE, "..", "data")
SRC = os.path.join(DATA, "pvs_flujos.xlsx")
TP = {"SUB_SOBERANO": "usd", "SUB_SOBERANO_DOLLAR_LINKED": "dl", "SUB_SOBERANO_TAMAR": "tamar", "SUB_SOBERANO_BADLAR": "badlar",
      "SUB_SOBERANO_CER": "cer", "SUB_SOBERANO_FIXED_RATE": "fija"}
# especie en pesos -> (MEP, cable); los que no tienen especie en dólares se valúan con el precio en pesos ÷ MEP
SPEC = {"BA37D": ("BA7DD", "BA7DC"), "BB37D": ("BB7DD", "BB7DC"), "BC37D": ("BC7DD", "BC7DC"), "BDC33": ("BC33D", "BC33C"),
        "BDC36": ("BC36D", "BC36C"), "CO26": ("CO26D", None), "CO32": ("CO32D", "CO32C"), "CO35": ("CO35D", "CO35C"),
        "ERM33": ("EM33D", "EM33C"), "NDG34": ("NDG4D", "NDG4C"), "NDT11": ("NDT1D", "NDT1C"), "NDT25": ("NDT5D", "NDT5C"),
        "PMM29": ("PM29D", "PM29C"), "PUA36": ("PUA6D", "PUA6C"), "PUL26": ("PU26D", "PU26C"), "SA24D": ("S24DD", "S24DC"),
        "SFD34": ("SFD4D", "SFD4C"), "SJO35": ("SJO5D", "SJO5C")}
EMIS = {"BA37D": "Buenos Aires", "BB37D": "Buenos Aires", "BC37D": "Buenos Aires", "CH24D": "Chubut",
        "CO24D": "Córdoba", "CO26": "Córdoba", "CO27D": "Córdoba", "CO32": "Córdoba", "CO35": "Córdoba", "CO2D7": "Córdoba", "COD7": "Córdoba", "COY27": "Córdoba",
        "ERF25": "Entre Ríos", "ERM33": "Entre Ríos", "ERE27": "Entre Ríos", "JUS22": "Jujuy",
        "NDT11": "Neuquén", "NDT25": "Neuquén", "NDG34": "Neuquén", "PMM29": "Mendoza", "PMA28": "Mendoza", "PMD26": "Mendoza", "PMM27": "Mendoza",
        "SA24D": "Salta", "SF27D": "Santa Fe", "SFD34": "Santa Fe", "SFN27": "Santa Fe", "TFU27": "Tierra del Fuego",
        "BDC28": "Ciudad de Buenos Aires", "BDC33": "Ciudad de Buenos Aires", "BDC36": "Ciudad de Buenos Aires",
        "BAF27": "Municipalidad de Córdoba", "BAM27": "Municipalidad de Córdoba", "BAO27": "Municipalidad de Córdoba", "BAS27": "Municipalidad de Córdoba",
        "BBD26": "Buenos Aires", "PBA27": "Buenos Aires", "PBA28": "Buenos Aires", "PBM27": "Buenos Aires", "PVC7O": "Banco Provincia de Buenos Aires",
        "BGD26": "Río Negro", "RND25": "Río Negro", "BRA27": "Municipalidad de Río Cuarto", "RIF25": "La Rioja (bono verde)", "SJO35": "San Juan",
        "BU2M7": "Chubut", "BU2N6": "Chubut", "BU3N6": "Chubut", "BU4N6": "Chubut", "BU5N6": "Chubut", "BUM27": "Chubut", "BUN26": "Chubut", "PUL26": "Chubut", "PUA36": "Chubut",
        "PR17": "Estado Nacional (Bono de Consolidación)"}


def main():
    hoy = os.environ.get("HOY") or datetime.date.today().isoformat()
    wb = openpyxl.load_workbook(SRC, read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    ix = {str(h): i for i, h in enumerate(rows[0])}
    by = {}
    for r in rows[1:]:
        if r[ix["ticker"]]:
            by.setdefault(r[ix["ticker"]], []).append(r)
    out = {}
    for t, rs in by.items():
        rs.sort(key=lambda r: r[ix["payment_date"]])
        mat = rs[-1][ix["payment_date"]]
        if mat < hoy:
            continue  # ya venció
        f = lambda r, c: float(r[ix[c]] or 0)
        fut = [r for r in rs if r[ix["payment_date"]] > hoy]
        past = [r for r in rs if r[ix["payment_date"]] <= hoy]
        o = {"t": TP.get(rs[0][ix["type"]], "otro"), "o": t, "iss": rs[0][ix["issue_date"]], "mat": mat,
             "r": round(f(fut[0], "interest_rate") * 100, 4) if fut else None,
             "res": round(f(past[-1], "residual_value") if past else 100.0, 4),
             "lc": past[-1][ix["payment_date"]] if past else rs[0][ix["issue_date"]],
             "b": int(float(rs[0][ix["days_convention"]] or 365)),
             "fl": [[r[ix["payment_date"]], round(f(r, "cash_flow"), 6), round(f(r, "interest_amount"), 6)] for r in fut],
             "fp": [[r[ix["payment_date"]], round(f(r, "cash_flow"), 6), round(f(r, "interest_amount"), 6)] for r in past]}
        if t in SPEC:
            o["d"], o["c"] = SPEC[t]
        out[t] = o
    try:  # los que se cargan a mano (x: 1, solo aparecen si se los busca, como PR17) no salen de la planilla: se conservan
        out.update({k: v for k, v in json.load(open(os.path.join(DATA, "pvs.json")))["on"].items() if v.get("x")})
    except Exception:
        pass
    json.dump({"hoy": hoy, "on": out}, open(os.path.join(DATA, "pvs.json"), "w"), separators=(",", ":"), ensure_ascii=False)
    json.dump({k: v for k, v in EMIS.items() if k in out}, open(os.path.join(DATA, "pvs_emisores.json"), "w"), ensure_ascii=False, separators=(",", ":"))
    print(len(out), "bonos provinciales vigentes")
    if os.environ.get("SERIES"):
        cierre(out)


def cierre(out):
    """Foto del último cierre completo: [precio, TIR del cierre, monto, monto promedio de 20 ruedas, ruedas operadas de 20]."""
    wb = openpyxl.load_workbook(os.environ["SERIES"], read_only=True)
    rows = list(wb.worksheets[0].iter_rows(values_only=True))
    ix = {str(h): i for i, h in enumerate(rows[0])}
    syms = set()
    for k, o in out.items():
        syms.add(k); syms.update(x for x in (o.get("d"), o.get("c")) if x)
    by = {}
    for r in rows[1:]:
        if r[ix["ticker"]] in syms and r[ix["closing_price"]]:
            by.setdefault(r[ix["ticker"]], {})[r[ix["date"]]] = r
    dates = sorted({d for m in by.values() for d in m})
    last = os.environ.get("CIERRE") or dates[-1]
    d20 = [d for d in dates if d <= last][-20:]
    tir = {}
    if os.environ.get("YIELDS"):
        for r in csv.DictReader(open(os.environ["YIELDS"])):
            if r["date"] == last and r["tir"]:
                tir[r["ticker"]] = float(r["tir"])
    px = {}
    for s, m in by.items():
        ds = [d for d in m if d <= last]
        if not ds:
            continue
        r = m[max(ds)]
        mo = lambda d: float(m[d][ix["effective_volume"]] or 0) if d in m else 0.0
        px[s] = [float(r[ix["closing_price"]]), tir.get(s), round(mo(last)), round(sum(mo(d) for d in d20) / len(d20)), sum(1 for d in d20 if mo(d) > 0)]
    json.dump({"d": last, "px": px}, open(os.path.join(DATA, "pvs_cierre.json"), "w"), separators=(",", ":"))
    print("cierre", last, len(px), "especies")


if __name__ == "__main__":
    main()
