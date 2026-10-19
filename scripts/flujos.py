"""Flujos de fondos de los bonos: la planilla data/flujos.xlsx es la fuente; la página los lleva adentro (PX.S, PSD.INST y PSOLD.INST).

    python3 scripts/flujos.py build    # planilla -> index.html y preview.html (lo corre solo el Action cuando cambia la planilla)
    python3 scripts/flujos.py export   # index.html -> planilla (para rearmarla desde lo que tiene la página)
    python3 scripts/flujos.py check    # verifica que la planilla y la página digan lo mismo

Hojas de la planilla (una fila por pago; los datos del bono se repiten en cada fila):
- «Pesos» y «Pesos vencidos»: Ticker, Clase, Familia, Emisión, Fecha de pago, Pago, Amortización, Tasa, Interés, Residual, Tipo.
- «Dólares»: Ticker, Emisión, Fecha de pago, Cupón, Amortización, Residual.
- «Precio de emisión» (opcional): Ticker, Precio. Para un bono nuevo sin historia de precios: arranca su serie con ese precio.
Fechas como fecha de Excel (o texto AAAA-MM-DD); montos por 100 de valor nominal; tasas en decimal (0,021 = 2,1%).
"""
import datetime as dt
import json
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX = os.path.join(ROOT, "data", "flujos.xlsx")
PAGES = [os.path.join(ROOT, f) for f in ("index.html", "preview.html")]
TICKERS = os.path.join(ROOT, "scripts", "pesos_tickers.json")
EPOCH = dt.date(1970, 1, 1)

H_PESOS = ["Ticker", "Clase", "Familia", "Emisión", "Fecha de pago", "Pago", "Amortización", "Tasa", "Interés", "Residual", "Tipo"]
H_USD = ["Ticker", "Emisión", "Fecha de pago", "Cupón", "Amortización", "Residual"]
CLASES = {"cap": "Lecap / Boncap", "cpn": "Bono a tasa fija con cupón", "cer0": "Lecer / Boncer cupón cero", "cerc": "Boncer con cupón",
          "tamar": "Bono TAMAR", "dl": "Dollar linked", "dual": "Dual (el mayor de dos pagos)", "dualtf": "Dual tasa fija / TAMAR"}
LEEME = [
    ("Flujos de fondos de los bonos de Bonerío Local", ""),
    ("", ""),
    ("Cómo se usa", "Esta planilla es la fuente de los flujos de la página. Para sumar o corregir un bono: editar las filas, guardar y subir el archivo a GitHub (data/flujos.xlsx, «Add file → Upload files» reemplaza el anterior). El Action «Flujos» actualiza la página solo en un par de minutos."),
    ("Una fila por pago", "Los datos del bono (Ticker, Clase, Familia, Emisión) se repiten en cada fila de sus pagos. El orden de las filas no importa salvo dentro de cada bono (por fecha)."),
    ("Montos", "Por 100 de valor nominal original. Tasas en decimal: 0,021 = 2,1%."),
    ("", ""),
    ("Hoja «Pesos»", ""),
    ("Clase", "cap = Lecap / Boncap (capitaliza y paga todo al vencimiento) · cpn = tasa fija con cupón · cer0 = CER cupón cero · cerc = Boncer con cupón · tamar = bono TAMAR · dl = dollar linked · dual = dual (paga el mayor de dos patas: dos filas con la misma fecha) · dualtf = dual tasa fija / TAMAR."),
    ("Familia", "fija, cer, tamar, dl, o las dos patas de un dual: cer+tamar, dl+tamar."),
    ("Pago", "Lo que paga en esa fecha (cupón + amortización). En una Lecap, el monto final por 100 VN. En los CER, TAMAR y dollar linked, en valores sin ajustar."),
    ("Amortización", "Capital que devuelve en esa fecha."),
    ("Tasa", "La del bono en decimal: TEM de la Lecap, cupón anual del CER, margen sobre TAMAR, etc."),
    ("Interés", "Interés de esa fecha (0 si capitaliza)."),
    ("Residual", "Valor nominal que queda después del pago."),
    ("Tipo", "FIXED_RATE, CER, TAMAR o DOLLAR_LINKED (en un dual, una fila de cada tipo)."),
    ("", ""),
    ("Hoja «Dólares»", "Globales y Bonares: Cupón y Amortización de cada fecha por 100 VN, y Residual (VN que queda)."),
    ("Hoja «Pesos vencidos»", "Bonos que ya vencieron: solo para mostrar su historia."),
    ("Hoja «Precio de emisión»", "Opcional. Para un bono nuevo que todavía no tiene cierres guardados: su precio de emisión, para arrancar la serie."),
]


def d2n(v):
    if isinstance(v, dt.datetime):
        v = v.date()
    if isinstance(v, dt.date):
        return (v - EPOCH).days
    if isinstance(v, (int, float)):
        return int(v)
    return (dt.date.fromisoformat(str(v).strip()[:10]) - EPOCH).days


def n2d(n):
    return EPOCH + dt.timedelta(days=int(n))


def num(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, str):
        v = v.strip().replace(",", ".")
    return float(v)


def read_page(path):
    lines = open(path, encoding="utf-8").read().split("\n")
    out = {}
    for name in ("PX", "PSD", "PSOLD"):
        pre = "const %s = " % name
        n = next(i for i, l in enumerate(lines) if l.startswith(pre))
        s = lines[n]; i = s.index("{"); j = s.rindex("}")
        out[name] = (n, s[:i], json.loads(s[i:j + 1]), s[j + 1:])
    return lines, out


def dump(o):
    return json.dumps(o, separators=(",", ":"), ensure_ascii=False)


# ---------- página -> planilla ----------
def export():
    _, P = read_page(PAGES[0])
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "Leeme"
    for a, b in LEEME:
        ws.append([a, b])
    ws["A1"].font = Font(bold=True, size=14)
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 1).font = Font(bold=True)
        ws.cell(r, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(r, 1).alignment = Alignment(vertical="top")
    ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 120

    def sheet(title, head, rows, fmt):
        w = wb.create_sheet(title); w.append(head)
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="44403C")
        for r in rows:
            w.append(r)
        for col, f in fmt.items():
            for r in range(2, w.max_row + 1):
                w.cell(r, col).number_format = f
        for i, h in enumerate(head, 1):
            w.column_dimensions[get_column_letter(i)].width = max(12, len(h) + 4)
        w.freeze_panes = "B2"; w.auto_filter.ref = w.dimensions
        return w

    def pesos_rows(inst):
        return [[t, b["k"], b["f"], n2d(b["iss"]), n2d(r[0]), r[1], r[2], r[3], r[4], r[5], r[6]] for t, b in inst.items() for r in b["r"]]
    fp = {4: "yyyy-mm-dd", 5: "yyyy-mm-dd", 6: "0.000000", 7: "0.000000", 8: "0.0000", 9: "0.000000", 10: "0.00"}
    sheet("Pesos", H_PESOS, pesos_rows(P["PSD"][2]["INST"]), fp)
    sheet("Dólares", H_USD, [[t, n2d(b["iss"]), n2d(r[0]), r[1], r[2], r[3]] for t, b in P["PX"][2]["S"].items() for r in b["r"]],
          {2: "yyyy-mm-dd", 3: "yyyy-mm-dd", 4: "0.0000", 5: "0.0000", 6: "0.00"})
    sheet("Pesos vencidos", H_PESOS, pesos_rows(P["PSOLD"][2]["INST"]), fp)
    sheet("Precio de emisión", ["Ticker", "Precio"], [], {2: "0.00"})
    wb.save(XLSX)
    print("planilla guardada:", os.path.relpath(XLSX, ROOT))


# ---------- planilla -> página ----------
def read_xlsx():
    wb = openpyxl.load_workbook(XLSX, data_only=True)

    def rows(name, head):
        if name not in wb.sheetnames:
            return []
        w = wb[name]; it = w.iter_rows(values_only=True); h = [str(x or "").strip() for x in next(it)]
        miss = [c for c in head if c not in h]
        if miss:
            sys.exit("hoja «%s»: faltan las columnas %s" % (name, ", ".join(miss)))
        ix = [h.index(c) for c in head]
        out = []
        for n, r in enumerate(it, 2):
            if r is None or all(v in (None, "") for v in r):
                continue
            out.append((n, [r[i] if i < len(r) else None for i in ix]))
        return out

    def pesos(name):
        inst = {}
        for n, (t, k, f, iss, fe, pago, am, tasa, intr, res, tipo) in rows(name, H_PESOS):
            try:
                t = str(t).strip(); b = inst.setdefault(t, {"k": str(k).strip(), "f": str(f).strip(), "iss": d2n(iss), "r": []})
                b["r"].append([d2n(fe), num(pago), num(am), num(tasa), num(intr), num(res), str(tipo).strip()])
            except Exception as e:
                sys.exit("hoja «%s», fila %d: %s" % (name, n, e))
        for b in inst.values():
            b["r"].sort(key=lambda r: r[0])
        return inst

    usd = {}
    for n, (t, iss, fe, cu, am, res) in rows("Dólares", H_USD):
        try:
            t = str(t).strip(); b = usd.setdefault(t, {"iss": d2n(iss), "r": []})
            b["r"].append([d2n(fe), num(cu), num(am), num(res)])
        except Exception as e:
            sys.exit("hoja «Dólares», fila %d: %s" % (n, e))
    for b in usd.values():
        b["r"].sort(key=lambda r: r[0])
    px = {str(t).strip(): num(p) for _, (t, p) in rows("Precio de emisión", ["Ticker", "Precio"]) if t and p not in (None, "")}
    return pesos("Pesos"), usd, pesos("Pesos vencidos"), px


def validate(inst, usd):
    err = []
    for t, b in inst.items():
        if b["k"] not in CLASES:
            err.append("%s: clase «%s» desconocida" % (t, b["k"]))
        for r in b["r"]:
            if r[6] not in ("FIXED_RATE", "CER", "TAMAR", "DOLLAR_LINKED"):
                err.append("%s: tipo «%s» desconocido" % (t, r[6]))
            if r[0] <= b["iss"]:
                err.append("%s: pago del %s anterior a la emisión" % (t, n2d(r[0])))
    for t, b in usd.items():
        for r in b["r"]:
            if r[0] <= b["iss"]:
                err.append("%s: pago del %s anterior a la emisión" % (t, n2d(r[0])))
    if err:
        sys.exit("la planilla tiene errores:\n  " + "\n  ".join(err))


def build(write=True):
    inst, usd, old, px = read_xlsx()
    validate(inst, usd)
    changed = []
    for path in PAGES:
        lines, P = read_page(path)
        new = {}
        n, a, o, b = P["PX"]; o["S"] = usd; new[n] = a + dump(o) + b
        n, a, o, b = P["PSD"]; o["INST"] = inst
        for t, ins in inst.items():  # bono nuevo sin historia: la serie arranca en la emisión, con su precio
            if t not in o["H"] and t in px:
                o["H"][t] = {"s": ins["iss"], "d": [0], "p": [px[t]]}
        new[n] = a + dump(o) + b
        n, a, o, b = P["PSOLD"]; o["INST"] = old; new[n] = a + dump(o) + b
        if any(lines[k] != v for k, v in new.items()):
            changed.append(os.path.basename(path))
            for k, v in new.items():
                lines[k] = v
            if write:
                open(path, "w", encoding="utf-8").write("\n".join(lines))
    sin_h = [t for t in inst if t not in P["PSD"][2]["H"]]
    if sin_h:
        print("ojo: sin historia de precios (cargar su precio en la hoja «Precio de emisión»):", ", ".join(sin_h))
    raw = open(TICKERS, encoding="utf-8").read(); tk = json.loads(raw); END = "\n" if raw.endswith("\n") else ""
    lst = tk if isinstance(tk, list) else None
    if lst is not None:
        falt = [t for t in inst if t not in lst]
        if falt and write:
            open(TICKERS, "w", encoding="utf-8").write(json.dumps(sorted(lst + falt), ensure_ascii=False) + END)
            print("tickers sumados a scripts/pesos_tickers.json:", ", ".join(falt))
    print("páginas actualizadas:" if write else "difieren de la planilla:", ", ".join(changed) or "ninguna")
    return changed


if __name__ == "__main__":
    cmd = sys.argv[1] if len(sys.argv) > 1 else ""
    if cmd == "export":
        export()
    elif cmd == "build":
        build()
    elif cmd == "check":
        sys.exit(1 if build(write=False) else 0)
    else:
        sys.exit(__doc__)
